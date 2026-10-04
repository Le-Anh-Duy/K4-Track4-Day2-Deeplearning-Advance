"""experiments.py - danh sách thí nghiệm, suy luận nhiều view, dự đoán chung kết, tổng hợp results.xlsx.

Notebook chỉ gọi các hàm ở đây và ở train.py; mọi huấn luyện vẫn đi qua train.run(Config).
"""
from __future__ import annotations

import json
import shutil
from dataclasses import asdict
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from torchvision.transforms.v2 import functional as TF

import train  # noqa: F401  (thêm repo root vào sys.path để import eval)
from dataset import CLASS_NAMES, NATIVE_SIZE, build_transforms, load_split, make_loader
from eval import load_group, mean_std, save_predictions
from inference import (aggregate_views, apply_temperature, fit_temperature, predict_logits,
                       view_hflip, views_multiscale)
from train import Config, curves_path, load_model, pred_path, run_dir

# --- Bước 1: >= 5 backbone (ResNet, ConvNeXt, 2 transformer, 2 mạng nhẹ), cùng công thức nền T00 ---
BACKBONES = [
    ("B01", "resnet50", "resnet50"),
    ("B02", "convnext_tiny", "convnext_tiny"),
    ("B03", "deit_small_patch16_224", "deit_small"),
    ("B04", "swin_tiny_patch4_window7_224", "swin_tiny"),
    ("B05", "efficientnet_b0", "efficientnet_b0"),
    ("B06", "mobilenetv3_large_100", "mobilenetv3"),
]

# --- Bước 2: mỗi dòng khác T00 đúng MỘT yếu tố: (exp_id, trục, mô tả, desc cho tên file, thay đổi) ---
TRAINING = [
    ("T00", "-", "công thức nền", "baseline", {}),
    ("T01", "A", "khởi tạo: từ đầu (không tiền huấn luyện)", "scratch", {"init": "scratch"}),
    ("T02", "A", "khởi tạo: đóng băng backbone, chỉ train head", "frozen", {"init": "frozen"}),
    ("T03", "B", "augmentation: + ColorJitter", "color", {"aug": "color"}),
    ("T04", "B", "augmentation: + TrivialAugmentWide", "trivial", {"aug": "trivial"}),
    ("T05", "B", "augmentation: + CutMix (alpha=1)", "cutmix", {"mix": "cutmix", "mix_alpha": 1.0}),
    ("T06", "C", "loss: label smoothing eps=0.1", "ls", {"loss": "ls", "label_smoothing": 0.1}),
    ("T07", "C", "loss: focal gamma=2", "focal", {"loss": "focal", "focal_gamma": 2.0}),
    ("T08", "C", "loss: CE trọng số lớp 1/n_c", "ce_weighted", {"loss": "ce_weighted", "class_weight_beta": 0.0}),
    ("T09", "D", "sampler cân bằng lớp", "balanced", {"sampler": "balanced"}),
    ("T10", "F", "EMA trọng số d=0.998", "ema", {"ema_decay": 0.998}),
]

SCALES = (224, 256, 288)
VIEW_METHODS = ("1view", "hflip", "5crop", "5crop_flip", "scale")


def method_sizes(method: str, img_size: int = 224) -> list[int]:
    """Độ phân giải của từng lượt forward trong một phương pháp (để đo độ trễ đúng chi phí)."""
    if method.startswith("res"):
        return [int(method[3:])]
    return {"1view": [img_size], "hflip": [img_size] * 2, "5crop": [img_size] * 5,
            "5crop_flip": [img_size] * 10, "scale": list(SCALES)}[method]


def run_views(model, df: pd.DataFrame, cfg: Config, method: str, device="cuda"):
    """Logit của mọi view của một phương pháp suy luận -> (filenames, y, [logits_view1, ...]).

    1view/hflip: Resize 256 + CenterCrop 224 (giống val). 5crop/5crop_flip/scale: ảnh gốc 256 rồi cắt/resize
    trên GPU. resS (ví dụ res288): resize cả ảnh về S (dò độ phân giải, FixRes).
    """
    if method.startswith("res"):
        tf = build_transforms(False, int(method[3:]), crop=False)
    elif method in ("1view", "hflip"):
        tf = build_transforms(False, cfg.img_size)
    else:
        tf = build_transforms(False, NATIVE_SIZE, crop=False)
    c = cfg.img_size
    views = {
        "1view": [None], "hflip": [None, view_hflip],
        "5crop": [lambda x, i=i: TF.five_crop(x, [c, c])[i] for i in range(5)],
        "5crop_flip": [lambda x, i=i, f=f: (view_hflip if f else (lambda z: z))(TF.five_crop(x, [c, c])[i])
                       for f in (0, 1) for i in range(5)],
        "scale": [lambda x, s=s: views_multiscale(x, [s])[0] for s in SCALES],
    }.get(method, [None])
    loader = make_loader(df, cfg.images_dir, tf, cfg.batch_size * 2, False, None, cfg.num_workers, cfg.cache)
    out = []
    for v in views:
        names, y, logits = predict_logits(model, loader, device, view=v, amp=cfg.amp)
        out.append(logits)
    return names, y, out


def alias_run(src: Config, dst: Config) -> None:
    """Dùng lại một run đã xong có CÙNG cấu hình huấn luyện (chỉ khác exp_id/desc) thay vì train lại.

    Ví dụ T00 seed0 = B0x của backbone đã chọn; F01 seed0 = run tốt nhất ở Bước 2. Ghi chú trong summary.json.
    """
    s, d = run_dir(src), run_dir(dst)
    if (d / "summary.json").exists():
        return
    a, b = asdict(src), asdict(dst)
    diff = sorted(k for k in a if a[k] != b[k] and k not in ("exp_id", "desc"))
    if diff:
        raise ValueError(f"{src.exp_id} và {dst.exp_id} khác cấu hình ở {diff}: phải train riêng")
    if not (s / "summary.json").exists():
        raise FileNotFoundError(f"{s} chưa chạy xong")
    shutil.copytree(s, d, dirs_exist_ok=True)
    for f in ("test_logits.npy", "final.json"):
        (d / f).unlink(missing_ok=True)
    (d / "config.json").write_text(json.dumps(b, indent=2))
    summary = json.loads((s / "summary.json").read_text())
    summary.update(exp_id=dst.exp_id, desc=dst.desc, config=b, note=f"dùng lại run {src.exp_id} seed{src.seed}")
    (d / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False))
    shutil.copy(pred_path(src, "val"), pred_path(dst, "val"))
    curves_path(dst).parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(curves_path(src), curves_path(dst))


def final_predict(cfg: Config, method: str = "1view", space: str = "prob", temperature: bool = False,
                  device="cuda") -> dict:
    """Bước 4: ghi predictions/<exp_id>_seed<k>_{val,test}.csv cho checkpoint best.pt của cfg.

    T (nếu temperature=True) khớp trên VAL, rồi áp dụng sang test; ghi thêm <exp_id>uncal_seed<k>_test.csv
    (cùng dự đoán, chưa hiệu chuẩn) để `eval.py grade --uncal` chấm I4(a).
    Test chạy ĐÚNG MỘT LẦN: nếu file test đã có thì dừng, không tính lại.
    """
    test_path = pred_path(cfg, "test")
    if test_path.exists():
        print(f"{test_path} đã có: không chạy lại test.")
        return json.loads((run_dir(cfg) / "final.json").read_text())
    model = load_model(cfg, device)
    _, val_df, test_df = load_split(cfg.labels_dir, cfg.fold)
    vnames, yv, vl = run_views(model, val_df, cfg, method, device)
    zv = np.log(np.clip(aggregate_views(vl, space), 1e-12, None))  # log-xác suất đóng vai logit
    T = fit_temperature(zv, yv) if temperature else 1.0

    tnames, yt, tl = run_views(model, test_df, cfg, method, device)  # lượt test duy nhất
    zt = np.log(np.clip(aggregate_views(tl, space), 1e-12, None))
    save_predictions(pred_path(cfg, "val"), vnames, yv, apply_temperature(zv, T))
    if temperature:
        save_predictions(Path(cfg.pred_dir) / f"{cfg.exp_id}uncal_seed{cfg.seed}_test.csv", tnames, yt,
                         apply_temperature(zt, 1.0))
    save_predictions(test_path, tnames, yt, apply_temperature(zt, T))
    info = {"exp_id": cfg.exp_id, "seed": cfg.seed, "method": method, "space": space, "temperature": T}
    (run_dir(cfg) / "final.json").write_text(json.dumps(info, indent=2))
    del model
    torch.cuda.empty_cache()
    return info


# --------------------------------------------------------------------------- #
# Tổng hợp bảng
# --------------------------------------------------------------------------- #
CONFIG_COLS = ("init", "aug", "mix", "loss", "label_smoothing", "focal_gamma", "sampler", "ema_decay",
               "optimizer", "lr_backbone", "lr_head")


def collect(out_dir: str | Path) -> pd.DataFrame:
    """Gom mọi runs/<exp_id>/seed<k>/summary.json thành một bảng."""
    rows = []
    for f in sorted(Path(out_dir).glob("*/seed*/summary.json")):
        s = json.loads(f.read_text())
        cfg = s.pop("config")
        rows.append({**s, **{k: cfg.get(k) for k in CONFIG_COLS}})
    return pd.DataFrame(rows)


def final_tables(pred_dir: str | Path, exp_ids, test_csv: str, labels_csv: str | None = None):
    """Sheet Final (mỗi seed một dòng + dòng mean ± std) và PerClass, tính bằng eval.py từ file dự đoán."""
    final, per_class = [], []
    for exp in exp_ids:
        g = load_group(str(Path(pred_dir) / f"{exp}_seed*_test.csv"), test_csv)
        val_f1 = []
        for p, m in zip(g.preds, g.metrics):
            vf = Path(pred_dir) / f"{exp}_seed{p.seed}_val.csv"
            v = load_group(str(vf), None).metrics[0]["macro_f1"] if vf.exists() else np.nan
            val_f1.append(v)
            final.append({"exp_id": exp, "seed": p.seed, "macro_f1_val": v, "macro_f1_test": m["macro_f1"],
                          "top1_test": m["top1"], "balanced_acc_test": m["balanced_acc"], "ece_test": m["ece"]})
        mv, sv = mean_std(val_f1)
        final.append({"exp_id": exp, "seed": f"mean ± std ({len(g.preds)} seed)",
                      "macro_f1_val": f"{mv:.4f} ± {sv:.4f}",
                      **{k: f"{g.summary[s][0]:.4f} ± {g.summary[s][1]:.4f}" for k, s in
                         (("macro_f1_test", "macro_f1"), ("top1_test", "top1"),
                          ("balanced_acc_test", "balanced_acc"), ("ece_test", "ece"))}})
        for i, c in enumerate(CLASS_NAMES):
            per_class.append({"exp_id": exp, "class": c, "n_test": int(g.metrics[0]["support"][i]),
                              **{f"{k}_mean": g.summary[k][0][i] for k in ("precision", "recall", "f1")},
                              **{f"{k}_std": g.summary[k][1][i] for k in ("precision", "recall", "f1")}})
    return pd.DataFrame(final), pd.DataFrame(per_class)


def write_results_xlsx(path: str | Path, sheets: dict[str, pd.DataFrame], highlight: dict[str, str] | None = None):
    """Ghi các sheet; cố định hàng tiêu đề, số 4 chữ số thập phân, tô vàng dòng có giá trị lớn nhất của cột
    trong `highlight` ({sheet: cột})."""
    from openpyxl.styles import Font, PatternFill
    highlight = highlight or {}
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
                for c in col[1:]:
                    if isinstance(c.value, float):
                        c.number_format = "0.0000"
            col = highlight.get(name)
            vals = pd.to_numeric(df[col], errors="coerce").to_numpy(dtype=float) if col in df else None
            if vals is not None and np.isfinite(vals).any():
                r = int(np.nanargmax(vals)) + 2
                for c in ws[r]:
                    c.fill = PatternFill("solid", fgColor="FFF2CC")
    return path
